"""
parser_service.py
─────────────────
Extracts raw text from PDF and DOCX files.
- PDF  → PyMuPDF (fitz)
- DOCX → python-docx
"""

import io
import fitz  # PyMuPDF
from docx import Document


def parse_file(file_bytes: bytes, file_type: str) -> str:
    """
    Parse a resume file and return its full text content.

    Args:
        file_bytes: Raw binary content of the file.
        file_type:  'pdf' or 'docx'.

    Returns:
        Extracted plain text string.
    """
    file_type = file_type.lower().strip(".")
    if file_type == "pdf":
        return _parse_pdf(file_bytes)
    elif file_type == "docx":
        return _parse_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")


def _parse_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF using PyMuPDF."""
    text_parts = []
    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        for page in doc:
            text_parts.append(page.get_text("text"))
    return "\n".join(text_parts).strip()


def _parse_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX using python-docx."""
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also extract table text
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())
    return "\n".join(paragraphs).strip()
